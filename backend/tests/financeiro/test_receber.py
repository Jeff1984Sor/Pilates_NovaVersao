"""Contas a receber (API): lista com chips, ATRASADO derivado por fuso, baixa/estorno, avulso, recibo, exportações."""

import io
from datetime import UTC, date, datetime
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.models import Auditoria
from app.modules.financeiro import consultas, servico
from app.modules.financeiro.models import ContaBancaria, ContaReceber, MovimentoConta, StatusReceber

HOJE = date(2026, 10, 25)


def _fixar_agora(monkeypatch, quando: datetime):
    monkeypatch.setattr(servico, "agora", lambda: quando)
    monkeypatch.setattr(consultas, "agora", lambda: quando)


@pytest.fixture(autouse=True)
def _relogio(monkeypatch):
    """Fixa 'agora' em 25/10/2026 12:00 (Brasília)."""
    _fixar_agora(monkeypatch, datetime(2026, 10, 25, 15, 0, tzinfo=UTC))


def _conta(db, aluno_id, valor, venc, status=StatusReceber.ABERTO, pagamento=None, **kw):
    c = ContaReceber(aluno_id=aluno_id, valor=Decimal(str(valor)), vencimento=venc, status=status, pagamento=pagamento,
                     competencia=venc.strftime("%Y-%m"), **kw)
    db.add(c)
    db.flush()
    return c


@pytest.fixture
def contas(db, cenario):
    """Maria: a vencer 100 (10/11), atrasada 200 (10/10), paga 300. João: atrasada 50 (01/09), cancelada 400, vence hoje 10."""
    a, o = cenario.aluno.id, cenario.outro.id
    d = SimpleNamespace()
    d.futuro = _conta(db, a, 100, date(2026, 11, 10))
    d.atrasada = _conta(db, a, 200, date(2026, 10, 10))
    d.paga = _conta(db, a, 300, date(2026, 9, 10), StatusReceber.PAGO, date(2026, 9, 12), forma_pagamento="PIX")
    d.atrasada2 = _conta(db, o, 50, date(2026, 9, 1))
    d.cancelada = _conta(db, o, 400, date(2026, 10, 1), StatusReceber.CANCELADO)
    d.hoje = _conta(db, o, 10, HOJE)
    db.commit()
    return d


def _chips(r):
    return {c["situacao"]: (c["quantidade"], Decimal(c["valor"])) for c in r["chips"]}


def _acoes(db, acao):
    return list(db.scalars(select(Auditoria).where(Auditoria.acao == acao)))


# ------------------------------------------------------------------------------------------------ lista e chips
def test_lista_com_chips_de_situacao_contagem_e_valor(client, adm, contas):
    r = client.get("/api/financeiro/receber", headers=adm).json()
    assert r["hoje"] == "2026-10-25" and r["total"] == 6
    ch = _chips(r)
    assert ch["TODOS"] == (6, Decimal("1060"))
    assert ch["ABERTO"] == (2, Decimal("110"))  # a vencer 100 + vence hoje 10 (vencer hoje NÃO é atraso)
    assert ch["ATRASADO"] == (2, Decimal("250"))
    assert ch["PAGO"] == (1, Decimal("300"))
    assert ch["CANCELADO"] == (1, Decimal("400"))


def test_filtro_de_situacao_atrasado_e_derivado_e_chips_ignoram_a_propria_situacao(client, adm, contas):
    r = client.get("/api/financeiro/receber", params={"situacao": "ATRASADO"}, headers=adm).json()
    assert r["total"] == 2 and {i["situacao"] for i in r["itens"]} == {"ATRASADO"}
    assert {i["status"] for i in r["itens"]} == {"ABERTO"}  # no banco continua ABERTO: ATRASADO nunca é gravado
    assert {i["dias_atraso"] for i in r["itens"]} == {15, 54}
    assert _chips(r)["TODOS"][0] == 6


def test_atrasado_respeita_o_fuso_de_sao_paulo(client, adm, contas, monkeypatch):
    # 26/10 01:30 UTC ainda é 25/10 22:30 em São Paulo: a conta que vence 25/10 NÃO está atrasada
    _fixar_agora(monkeypatch, datetime(2026, 10, 26, 1, 30, tzinfo=UTC))
    r = client.get("/api/financeiro/receber", params={"situacao": "ATRASADO"}, headers=adm).json()
    assert contas.hoje.id not in [i["id"] for i in r["itens"]] and r["hoje"] == "2026-10-25"
    # 26/10 03:30 UTC = 26/10 00:30 em SP: agora ela venceu ontem
    _fixar_agora(monkeypatch, datetime(2026, 10, 26, 3, 30, tzinfo=UTC))
    r = client.get("/api/financeiro/receber", params={"situacao": "ATRASADO"}, headers=adm).json()
    assert contas.hoje.id in [i["id"] for i in r["itens"]] and r["hoje"] == "2026-10-26"
    assert _chips(r)["ATRASADO"] == (3, Decimal("260"))


def test_filtros_aluno_periodo_competencia_busca_e_plano(client, adm, db, cenario, contas):
    def total(**p):
        return client.get("/api/financeiro/receber", params=p, headers=adm).json()["total"]

    assert total(aluno_id=cenario.outro.id) == 3
    assert total(de="2026-10-01", ate="2026-10-31") == 3  # atrasada (10/10), cancelada (01/10) e vence hoje
    assert total(competencia="2026-09") == 2
    assert total(q="joão") == 3 and total(q="maria") == 3 and total(q="inexistente") == 0
    assert total(situacao="PAGO", aluno_id=cenario.aluno.id) == 1
    # plano: conta do contrato aparece no filtro de plano
    c = _conta(db, cenario.aluno.id, 300, date(2026, 12, 28), contrato_id=cenario.contrato.id)
    db.commit()
    r = client.get("/api/financeiro/receber", params={"plano_id": cenario.plano.id}, headers=adm).json()
    assert [i["id"] for i in r["itens"]] == [c.id]
    item = r["itens"][0]
    assert item["plano_nome"] == "Plano 2x" and item["subcategoria_nome"] == "Pilates mensal"
    assert item["origem"].startswith("Contrato #") and item["aluno_nome"] == "Maria Souza"


def test_paginacao_e_ordem_por_vencimento(client, adm, contas):
    r = client.get("/api/financeiro/receber", params={"por_pagina": 4, "pagina": 2}, headers=adm).json()
    assert len(r["itens"]) == 2 and r["total"] == 6 and r["pagina"] == 2
    venc = [i["vencimento"] for i in client.get("/api/financeiro/receber", headers=adm).json()["itens"]]
    assert venc == sorted(venc)


# ------------------------------------------------------------------------------------------------- baixa/estorno
def test_baixa_padrao_hoje_de_sao_paulo_e_auditoria(client, adm, db, contas):
    r = client.post(f"/api/financeiro/receber/{contas.atrasada.id}/baixa", json={"forma_pagamento": "PIX"}, headers=adm)
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["status"] == "PAGO" and j["situacao"] == "PAGO" and j["pagamento"] == "2026-10-25" and j["forma_pagamento"] == "PIX"
    a = _acoes(db, "baixar")
    assert len(a) == 1 and a[0].entidade == "conta_receber" and a[0].entidade_id == str(contas.atrasada.id)
    assert a[0].usuario_nome


def test_baixa_com_data_e_nao_aceita_futura_nem_repetida_nem_cancelada(client, adm, contas):
    url = f"/api/financeiro/receber/{contas.futuro.id}/baixa"
    assert client.post(url, json={"data_pagamento": "2026-11-01"}, headers=adm).status_code == 422
    r = client.post(url, json={"data_pagamento": "2026-10-20", "forma_pagamento": "DINHEIRO"}, headers=adm)
    assert r.json()["pagamento"] == "2026-10-20"
    assert client.post(url, json={}, headers=adm).status_code == 409  # já paga
    assert client.post(f"/api/financeiro/receber/{contas.cancelada.id}/baixa", json={}, headers=adm).status_code == 409
    assert client.post(f"/api/financeiro/receber/{contas.atrasada.id}/baixa", json={"forma_pagamento": "CHEQUE"},
                       headers=adm).status_code == 422


def test_baixa_com_conta_bancaria_gera_entrada_e_estorno_remove(client, adm, db, contas):
    banco = ContaBancaria(banco="Banco X", agencia="1", conta="2", saldo_inicial=Decimal("1000"))
    db.add(banco)
    db.commit()
    r = client.post(f"/api/financeiro/receber/{contas.atrasada.id}/baixa", json={"conta_bancaria_id": banco.id}, headers=adm)
    assert r.status_code == 200
    movs = db.scalars(select(MovimentoConta)).all()
    assert len(movs) == 1 and movs[0].tipo.value == "ENTRADA" and movs[0].valor == Decimal("200.00")
    assert movs[0].conta_receber_id == contas.atrasada.id and movs[0].data == HOJE
    assert Decimal(client.get("/api/financeiro/bancos/contas", headers=adm).json()[0]["saldo_atual"]) == Decimal("1200.00")

    e = client.delete(f"/api/financeiro/receber/{contas.atrasada.id}/baixa", headers=adm)
    assert e.status_code == 200 and e.json()["status"] == "ABERTO" and e.json()["pagamento"] is None
    db.expire_all()
    assert db.scalars(select(MovimentoConta)).all() == []
    assert Decimal(client.get("/api/financeiro/bancos/contas", headers=adm).json()[0]["saldo_atual"]) == Decimal("1000.00")
    assert len(_acoes(db, "estornar_baixa")) == 1
    assert client.delete(f"/api/financeiro/receber/{contas.atrasada.id}/baixa", headers=adm).status_code == 409


def test_baixa_em_conta_bancaria_inativa_ou_inexistente(client, adm, db, contas):
    banco = ContaBancaria(banco="B", agencia="1", conta="2", ativo=False)
    db.add(banco)
    db.commit()
    url = f"/api/financeiro/receber/{contas.futuro.id}/baixa"
    assert client.post(url, json={"conta_bancaria_id": banco.id}, headers=adm).status_code == 422
    assert client.post(url, json={"conta_bancaria_id": 9999}, headers=adm).status_code == 404
    db.expire_all()
    assert db.get(ContaReceber, contas.futuro.id).status == StatusReceber.ABERTO  # nada ficou pela metade


def test_estorno_recusado_com_recibo_saude_gerado(client, adm, db, contas):
    contas.paga.recibo_saude_gerado_em = datetime(2026, 10, 1, tzinfo=UTC)
    db.commit()
    r = client.delete(f"/api/financeiro/receber/{contas.paga.id}/baixa", headers=adm)
    assert r.status_code == 409 and "Receita Saúde" in r.json()["detail"]


# ------------------------------------------------------------------------------------ cancelar/editar/excluir
def test_cancelar_exige_motivo_e_reabrir(client, adm, db, contas):
    url = f"/api/financeiro/receber/{contas.futuro.id}"
    assert client.post(f"{url}/cancelar", json={}, headers=adm).status_code == 422
    r = client.post(f"{url}/cancelar", json={"motivo": "Aluno trancou o plano"}, headers=adm)
    assert r.status_code == 200 and r.json()["situacao"] == "CANCELADO"
    assert r.json()["motivo_cancelamento"] == "Aluno trancou o plano"
    assert client.post(f"{url}/cancelar", json={"motivo": "de novo"}, headers=adm).status_code == 409
    assert client.post(f"{url}/reabrir", headers=adm).json()["situacao"] == "ABERTO"
    assert client.post(f"/api/financeiro/receber/{contas.paga.id}/cancelar", json={"motivo": "teste"}, headers=adm).status_code == 409
    assert len(_acoes(db, "cancelar")) == 1 and len(_acoes(db, "reabrir")) == 1


def test_editar_conta_aberta_e_paga(client, adm, db, contas):
    r = client.put(f"/api/financeiro/receber/{contas.futuro.id}",
                   json={"valor": "120.50", "descricao": "Ajuste", "vencimento": "2026-11-15"}, headers=adm)
    assert r.status_code == 200 and r.json()["valor"] == "120.50" and r.json()["vencimento"] == "2026-11-15"
    assert client.put(f"/api/financeiro/receber/{contas.paga.id}", json={"valor": "1"}, headers=adm).status_code == 409
    assert client.put(f"/api/financeiro/receber/{contas.paga.id}", json={"descricao": "Obs"}, headers=adm).status_code == 200
    assert client.put(f"/api/financeiro/receber/{contas.cancelada.id}", json={"descricao": "x"}, headers=adm).status_code == 409
    assert client.put(f"/api/financeiro/receber/{contas.futuro.id}", json={"valor": "-1"}, headers=adm).status_code == 422
    assert len(_acoes(db, "editar")) == 2


def test_excluir_aberta_ok_e_paga_recusada(client, adm, db, contas):
    assert client.delete(f"/api/financeiro/receber/{contas.paga.id}", headers=adm).status_code == 409
    assert client.delete(f"/api/financeiro/receber/{contas.futuro.id}", headers=adm).status_code == 200
    assert client.get(f"/api/financeiro/receber/{contas.futuro.id}", headers=adm).status_code == 404
    assert len(_acoes(db, "excluir")) == 1


# ----------------------------------------------------------------------------------------------------- avulso
def test_lancar_valor_avulso(client, adm, db, cenario):
    corpo = {"aluno_id": cenario.aluno.id, "descricao": "Taxa de matrícula", "valor": "150.00", "vencimento": "2026-11-05"}
    r = client.post("/api/financeiro/receber/avulso", json=corpo, headers=adm)
    assert r.status_code == 201, r.text
    j = r.json()
    assert j["origem"] == "Taxa de matrícula" and j["competencia"] == "2026-11" and j["contrato_id"] is None
    assert j["situacao"] == "ABERTO"
    assert len(_acoes(db, "lancar_avulso")) == 1
    assert client.post("/api/financeiro/receber/avulso", json={**corpo, "valor": "0"}, headers=adm).status_code == 422
    assert client.post("/api/financeiro/receber/avulso", json={**corpo, "aluno_id": 9999}, headers=adm).status_code == 404


# ----------------------------------------------------------------------------------------------------- recibo
def test_recibo_pdf_so_para_conta_paga(client, adm, db, contas):
    r = client.get(f"/api/financeiro/receber/{contas.paga.id}/recibo", headers=adm)
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf" and r.content.startswith(b"%PDF")
    assert f"recibo-{contas.paga.id}.pdf" in r.headers["content-disposition"]
    assert client.get(f"/api/financeiro/receber/{contas.futuro.id}/recibo", headers=adm).status_code == 409
    assert len(_acoes(db, "emitir_recibo")) == 1


def test_recibo_com_logo_do_estudio(client, adm, db, contas):
    """Identidade do banco: nome e logo (PNG) entram no PDF; logo ilegível não derruba o recibo."""
    from app.core import storage
    from app.core.models import Identidade

    png = bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d49444154789c6360f8cfc0f01f0005000201"
        "a0b9c8db0000000049454e44ae426082")
    caminho = storage.salvar("identidade", "logo.png", png, nome_base="logo", substituir=True)
    db.add(Identidade(id=1, nome_exibicao="Studio Mayris Pilates", logo_path=caminho))
    db.commit()
    r = client.get(f"/api/financeiro/receber/{contas.paga.id}/recibo", headers=adm)
    assert r.status_code == 200 and r.content.startswith(b"%PDF") and b"/Image" in r.content


# ------------------------------------------------------------------------------------------------- exportações
def test_exportar_excel_geral_e_por_aluno(client, adm, cenario, contas):
    r = client.get("/api/financeiro/receber/exportar", params={"formato": "xlsx"}, headers=adm)
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    linhas = list(load_workbook(io.BytesIO(r.content)).active.values)
    assert linhas[0][0] == "Aluno" and len(linhas) == 1 + 6 + 1  # cabeçalho + 6 faturas + total
    assert linhas[-1][-1] == 1060.0
    r = client.get("/api/financeiro/receber/exportar", params={"formato": "xlsx", "aluno_id": cenario.aluno.id}, headers=adm)
    linhas = list(load_workbook(io.BytesIO(r.content)).active.values)
    assert len(linhas) == 1 + 3 + 1 and linhas[-1][-1] == 600.0
    assert "faturas-aluno-" in r.headers["content-disposition"]


def test_exportar_pdf_com_filtro(client, adm, contas):
    r = client.get("/api/financeiro/receber/exportar", params={"formato": "pdf", "situacao": "ATRASADO"}, headers=adm)
    assert r.status_code == 200 and r.content.startswith(b"%PDF") and "faturas-geral.pdf" in r.headers["content-disposition"]


# -------------------------------------------------------------------------------------------- WhatsApp (extensão)
def test_cobranca_whatsapp_sem_cobrador_responde_503_e_com_cobrador_funciona(client, adm, contas, db):
    url = f"/api/financeiro/receber/{contas.atrasada.id}/cobrar-whatsapp"
    cobrador_real = servico._cobrador  # o módulo de comunicação o registra ao carregar; aqui simulamos "sem WhatsApp"
    servico.definir_cobrador(None)
    try:
        assert client.get("/api/financeiro/receber/cobranca-whatsapp/status", headers=adm).json()["enviado"] is False
        r = client.post(url, headers=adm)
        assert r.status_code == 503 and "comunicação" in r.json()["detail"]
    finally:
        servico.definir_cobrador(cobrador_real)

    servico.definir_cobrador(lambda db_, conta, aluno: f"Olá {aluno.nome}, fatura de R$ {conta.valor}")
    try:
        r = client.post(url, headers=adm)
        assert r.status_code == 200 and "Maria" in r.json()["mensagem"]
        assert client.post(f"/api/financeiro/receber/{contas.paga.id}/cobrar-whatsapp", headers=adm).status_code == 409
        assert len(_acoes(db, "cobrar_whatsapp")) == 1
    finally:
        servico.definir_cobrador(cobrador_real)


def test_resumo_do_aluno_na_api(client, adm, cenario, contas):
    r = client.get(f"/api/financeiro/receber/aluno/{cenario.aluno.id}/resumo", headers=adm).json()
    assert Decimal(r["total_atrasado"]) == 200 and Decimal(r["total_aberto"]) == 100 and Decimal(r["total_pago"]) == 300
    assert r["em_dia"] is False and r["proximo_vencimento"] == "2026-11-10"
    assert client.get("/api/financeiro/receber/aluno/9999/resumo", headers=adm).status_code == 404
