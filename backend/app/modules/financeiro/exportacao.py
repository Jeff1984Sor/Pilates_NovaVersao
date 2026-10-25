"""Exportações do financeiro: Excel (openpyxl) e PDF (reportlab), com a identidade do estúdio (nome e logo do banco)."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.core import storage
from app.core.db import agora
from app.core.models import Identidade
from app.domain.documentos import so_digitos

COR_CABECALHO = "#1d4d4d"
COR_ZEBRA = "#f6f7fb"
COR_GRADE = "#dbe4ea"
FUSO = ZoneInfo("America/Sao_Paulo")


def moeda(v: object) -> str:
    n = Decimal(str(v or 0)).quantize(Decimal("0.01"))
    s = f"{abs(n):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"-R$ {s}" if n < 0 else f"R$ {s}"


def data_br(d: date | datetime | None) -> str:
    if not d:
        return "-"
    if isinstance(d, datetime):
        d = d.astimezone(FUSO).date() if d.tzinfo else d.date()
    return d.strftime("%d/%m/%Y")


def cpf_fmt(cpf: str | None) -> str:
    d = so_digitos(cpf)
    return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}" if len(d) == 11 else (cpf or "-")


@dataclass
class Tabela:
    titulo: str
    colunas: list[str]
    linhas: list[list]  # Decimal vira moeda; date vira dd/mm/aaaa; resto, texto
    direita: set[int] = field(default_factory=set)  # índices de colunas alinhadas à direita
    rodape: list | None = None  # linha de totais


@dataclass
class Identidade_:
    nome: str
    logo: bytes | None


def identidade_do_estudio(db: Session) -> Identidade_:
    ident = db.get(Identidade, 1)
    nome = (ident.nome_exibicao if ident else None) or "Pilates"
    logo = None
    if ident and ident.logo_path:
        try:
            logo = storage.ler(ident.logo_path)
        except (FileNotFoundError, storage.ArquivoInvalido):
            logo = None
    return Identidade_(nome=nome, logo=logo)


# ------------------------------------------------------------------------------------------------------- Excel
def _celula_xlsx(v):
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, datetime):
        return v.astimezone(FUSO).replace(tzinfo=None) if v.tzinfo else v
    return v


def gerar_xlsx(tabelas: list[Tabela]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    usados: set[str] = set()
    for t in tabelas:
        nome = "".join(c for c in t.titulo if c not in '[]:*?/\\')[:31] or "Planilha"
        base, i = nome, 2
        while nome in usados:
            nome = f"{base[:28]}_{i}"
            i += 1
        usados.add(nome)
        ws = wb.create_sheet(nome)
        ws.append(t.colunas)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=COR_CABECALHO.lstrip("#"))
            c.alignment = Alignment(horizontal="center")
        for linha in t.linhas + ([t.rodape] if t.rodape else []):
            ws.append([_celula_xlsx(v) for v in linha])
        if t.rodape:
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = '"R$" #,##0.00'
                elif isinstance(c.value, datetime):
                    c.number_format = "dd/mm/yyyy hh:mm"
                elif isinstance(c.value, date):
                    c.number_format = "dd/mm/yyyy"
        for idx, col in enumerate(t.colunas, start=1):
            maior = max([len(str(col))] + [len(str(_texto(linha[idx - 1]))) for linha in t.linhas if idx - 1 < len(linha)] or [10])
            ws.column_dimensions[get_column_letter(idx)].width = min(max(maior + 2, 10), 60)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _texto(v) -> str:
    if v is None or v == "":
        return "-"
    if isinstance(v, Decimal):
        return moeda(v)
    if isinstance(v, datetime | date):
        return data_br(v)
    return str(v)


# --------------------------------------------------------------------------------------------------------- PDF
def _cabecalho(ident: Identidade_, titulo: str, subtitulo: str | None, estilos) -> list:
    nome = Paragraph(f"<b>{_esc(ident.nome)}</b>", estilos["Normal"])
    bloco = [nome, Paragraph(_esc(titulo), estilos["Title"])]
    if subtitulo:
        bloco.append(Paragraph(_esc(subtitulo), estilos["Normal"]))
    if ident.logo:
        try:
            img = Image(io.BytesIO(ident.logo))
            img._restrictSize(28 * mm, 18 * mm)
            return [Table([[img, bloco]], colWidths=[32 * mm, None], style=[("VALIGN", (0, 0), (-1, -1), "TOP")]), Spacer(1, 4)]
        except Exception:  # logo em formato que o reportlab não lê (ex.: WEBP): segue só com o nome
            pass
    return [*bloco, Spacer(1, 4)]


def _esc(txt: str) -> str:
    return (txt or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _tabela_pdf(t: Tabela, estilos, largura_util: float):
    pequeno = ParagraphStyle("cel", parent=estilos["Normal"], fontSize=8, leading=10)
    cab = ParagraphStyle("cab", parent=pequeno, textColor=colors.white, fontName="Helvetica-Bold")
    dados = [[Paragraph(_esc(c), cab) for c in t.colunas]]
    for linha in t.linhas + ([t.rodape] if t.rodape else []):
        dados.append([Paragraph(_esc(_texto(v)), pequeno) for v in linha])
    n = max(len(t.colunas), 1)
    tb = Table(dados, repeatRows=1, colWidths=[largura_util / n] * n)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(COR_CABECALHO)),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor(COR_GRADE)),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(COR_ZEBRA)]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    for i in t.direita:
        estilo.append(("ALIGN", (i, 1), (i, -1), "RIGHT"))
    if t.rodape:
        estilo.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#e8eef2")))
    tb.setStyle(TableStyle(estilo))
    return tb


def gerar_pdf(ident: Identidade_, titulo: str, tabelas: list[Tabela], *, subtitulo: str | None = None,
              paisagem: bool = True) -> bytes:
    buf = io.BytesIO()
    pagina = landscape(A4) if paisagem else A4
    doc = SimpleDocTemplate(buf, pagesize=pagina, title=titulo, leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm)
    estilos = getSampleStyleSheet()
    util = pagina[0] - 28 * mm
    el = _cabecalho(ident, titulo, subtitulo, estilos)
    for t in tabelas:
        if t.titulo and len(tabelas) > 1:
            el += [Spacer(1, 8), Paragraph(f"<b>{_esc(t.titulo)}</b>", estilos["Heading4"])]
        el.append(_tabela_pdf(t, estilos, util))
    gerado = agora().astimezone(FUSO).strftime("%d/%m/%Y %H:%M")
    el += [Spacer(1, 6), Paragraph(f"Gerado em {gerado}", ParagraphStyle("rod", parent=estilos["Normal"], fontSize=7,
                                                                       textColor=colors.grey))]
    doc.build(el)
    return buf.getvalue()


def gerar_recibo_pdf(ident: Identidade_, *, numero: int, aluno_nome: str, aluno_cpf: str | None, pagador_cpf: str | None,
                     origem: str, competencia: str | None, vencimento: date, pagamento: date | None, valor: Decimal,
                     forma_pagamento: str | None, descricao: str | None) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=f"Recibo {numero}", leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm)
    est = getSampleStyleSheet()
    el = _cabecalho(ident, "Recibo de Pagamento", f"Recibo nº {numero}", est)
    linhas = [
        ["Recebemos de", aluno_nome],
        ["CPF", cpf_fmt(aluno_cpf)],
    ]
    if pagador_cpf and so_digitos(pagador_cpf) != so_digitos(aluno_cpf):
        linhas.append(["CPF do pagador", cpf_fmt(pagador_cpf)])
    linhas += [
        ["Referente a", descricao or origem],
        ["Origem", origem],
        ["Competência", competencia or "-"],
        ["Vencimento", data_br(vencimento)],
        ["Data do pagamento", data_br(pagamento)],
        ["Forma de pagamento", (forma_pagamento or "-").capitalize()],
        ["Valor", moeda(valor)],
    ]
    tb = Table([[Paragraph(f"<b>{_esc(a)}</b>", est["Normal"]), Paragraph(_esc(b), est["Normal"])] for a, b in linhas],
               colWidths=[48 * mm, None])
    tb.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor(COR_GRADE)),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor(COR_ZEBRA)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    el += [Spacer(1, 10), tb, Spacer(1, 22),
           Paragraph(f"Pagamento recebido e quitado em {data_br(pagamento)}. Obrigado!", est["Normal"]),
           Spacer(1, 30), Paragraph("_" * 42, est["Normal"]), Paragraph(_esc(ident.nome), est["Normal"])]
    doc.build(el)
    return buf.getvalue()
